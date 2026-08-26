"""CSV/xlsx activities import — SPEC.md §4.3. Load-bearing, not a fallback:
everything the reconstruction knows about the pre-API era comes from here.

Two-phase: parse+preview first, commit on confirmation. Rows land in the same
`activities` table through the same multiset dedup as the API sync, so
re-importing or overlapping the API window is safe.
"""

import csv
import io
import json
import logging
import re

from sqlalchemy.ext.asyncio import async_sessionmaker

from .activities import assign_hashes, parse_date, parse_money_cents, parse_qty_micro, upsert_activities

log = logging.getLogger(__name__)

# Fuzzy header → field mapping for Questrade's web-app export.
HEADER_MAP = {
    "transaction date": "transaction_date",
    "trade date": "trade_date",
    "settlement date": "settlement_date",
    "action": "action",
    "symbol": "ticker",
    "description": "description",
    "quantity": "qty",
    "price": "price",
    "gross amount": "gross",
    "commission": "commission",
    "net amount": "net",
    "currency": "currency",
    "account #": "account_id",
    "account number": "account_id",
    "activity type": "type",
    "account type": "account_type",
}


def _match_header(h: str) -> str | None:
    key = re.sub(r"\s+", " ", str(h or "").strip().lower())
    if key in HEADER_MAP:
        return HEADER_MAP[key]
    for candidate, field in HEADER_MAP.items():
        if candidate in key:
            return field
    return None


def _read_table(filename: str, content: bytes) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def parse_export(filename: str, content: bytes) -> dict:
    """Parse the export into normalized activity dicts + a mapping report."""
    table = _read_table(filename, content)
    if not table:
        return {"rows": [], "errors": ["file is empty"], "header_map": {}, "unmapped_headers": []}

    # Find the header row (the export sometimes has preamble lines).
    header_idx, fields = None, None
    for i, row in enumerate(table[:10]):
        mapped = [_match_header(c) for c in row]
        if sum(1 for m in mapped if m) >= 4:
            header_idx, fields = i, mapped
            break
    if header_idx is None:
        return {"rows": [], "errors": ["could not locate a header row"], "header_map": {}, "unmapped_headers": []}

    header_map = {str(table[header_idx][i]): f for i, f in enumerate(fields) if f}
    unmapped = [str(c) for i, c in enumerate(table[header_idx]) if c and not fields[i]]

    rows, errors = [], []
    for line_no, raw in enumerate(table[header_idx + 1 :], start=header_idx + 2):
        if not any(c not in (None, "") for c in raw):
            continue
        rec = {}
        for i, field in enumerate(fields):
            if field and i < len(raw):
                rec[field] = raw[i]
        txn_date = parse_date(rec.get("transaction_date"))
        if txn_date is None:
            errors.append(f"line {line_no}: unparseable transaction date {rec.get('transaction_date')!r}")
            continue
        price_micro = parse_qty_micro(rec.get("price"))  # same 1e-6 scaling as qty
        ticker = str(rec.get("ticker") or "").strip().upper() or None
        rows.append({
            "account_id": str(rec.get("account_id") or "").strip(),
            "trade_date": parse_date(rec.get("trade_date")),
            "transaction_date": txn_date,
            "settlement_date": parse_date(rec.get("settlement_date")),
            "type": re.sub(r"\s+", " ", str(rec.get("type") or "").strip()).lower(),
            "action": re.sub(r"\s+", " ", str(rec.get("action") or "").strip()).lower(),
            "ticker": ticker,
            "qty_micro": parse_qty_micro(rec.get("qty")),
            "price_micro": price_micro if price_micro else None,
            "gross_amount_cents": parse_money_cents(rec.get("gross")),
            "commission_cents": parse_money_cents(rec.get("commission")),
            "net_amount_cents": parse_money_cents(rec.get("net")),
            "currency": str(rec.get("currency") or "").strip().upper() or None,
            "description": str(rec.get("description") or "").strip() or None,
            "raw_json": json.dumps({"source": "csv_import", "line": line_no}),
        })
    missing_account = sum(1 for r in rows if not r["account_id"])
    if missing_account:
        errors.append(f"{missing_account} rows have no account number and will import under account ''")
    return {"rows": rows, "errors": errors, "header_map": header_map, "unmapped_headers": unmapped}


def preview_report(parsed: dict) -> dict:
    rows = parsed["rows"]
    by_type: dict[str, int] = {}
    for r in rows:
        by_type[r["type"] or "?"] = by_type.get(r["type"] or "?", 0) + 1
    dates = [r["transaction_date"] for r in rows]
    return {
        "n_rows": len(rows),
        "by_type": by_type,
        "date_range": [min(dates).isoformat(), max(dates).isoformat()] if dates else None,
        "accounts": sorted({r["account_id"] for r in rows}),
        "header_map": parsed["header_map"],
        "unmapped_headers": parsed["unmapped_headers"],
        "errors": parsed["errors"],
        "sample": [
            {k: (v.isoformat() if hasattr(v, "isoformat") else v) for k, v in r.items() if k != "raw_json"}
            for r in rows[:10]
        ],
    }


async def commit_import(sessions: async_sessionmaker, parsed: dict) -> dict:
    rows = assign_hashes(parsed["rows"])
    new = await upsert_activities(sessions, rows)
    return {"imported": new, "duplicates": len(rows) - new, "errors": parsed["errors"]}
